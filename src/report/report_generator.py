# -*- coding: utf-8 -*-
"""Word 检测报告生成（PV-S3 语义分割版）。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches, Pt

from src.utils.config import PROJECT_ROOT, REPORTS_DIR, ensure_directories


def generate_report(result_dict: dict[str, Any]) -> str:
    """
    根据检测结果字典生成 Word 报告。
    文件名：report_图片名_时间.docx
    返回报告绝对路径。
    """
    ensure_directories()
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    original_path = Path(result_dict["original_path"])
    if not original_path.is_absolute():
        original_path = PROJECT_ROOT / original_path

    image_name = original_path.name
    stem = original_path.stem
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = REPORTS_DIR / f"report_{stem}_{ts}.docx"

    doc = Document()
    title = doc.add_heading(
        "基于 PV-S3 半监督语义分割的光伏板 EL 图像缺陷检测报告",
        level=0,
    )
    title.runs[0].font.size = Pt(16)

    doc.add_paragraph(
        "项目名称：基于 PV-S3 半监督语义分割的光伏板 EL 图像缺陷检测与智能运维辅助系统"
    )
    doc.add_paragraph(
        f"检测时间：{result_dict.get('detect_time', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}"
    )
    doc.add_paragraph(f"图片名称：{image_name}")
    doc.add_paragraph(
        f"推理模式：{result_dict.get('mode', 'unknown')} "
        f"（{result_dict.get('model_version', '')}）"
    )
    doc.add_paragraph(
        f"置信度阈值：{result_dict.get('confidence_threshold', 'N/A')}"
    )

    doc.add_heading("一、图像与分割结果", level=1)

    def add_image_if_exists(rel_key: str, caption: str) -> None:
        rel = result_dict.get(rel_key)
        if not rel:
            return
        p = Path(rel)
        if not p.is_absolute():
            p = PROJECT_ROOT / p
        if p.is_file():
            doc.add_paragraph(caption)
            doc.add_picture(str(p), width=Inches(5.5))

    add_image_if_exists("original_path", "原图：")
    add_image_if_exists("colorized_mask_path", "分类预测图（5类语义分割）：")
    add_image_if_exists("heatmap_path", "置信度热图：")
    add_image_if_exists("mask_path", "缺陷区域二值 Mask（置信度阈值过滤后）：")
    add_image_if_exists("overlay_path", "原图与分类预测叠加：")

    # ---- 缺陷类别分布 ----
    doc.add_heading("二、缺陷类别分布", level=1)
    per_class = result_dict.get("per_class_areas") or {}
    if per_class:
        class_table = doc.add_table(rows=1, cols=3)
        class_table.style = "Table Grid"
        hdr = class_table.rows[0].cells
        hdr[0].text = "缺陷类别"
        hdr[1].text = "像素面积"
        hdr[2].text = "占比（总缺陷中）"
        total_defect = sum(per_class.values())
        for cls_name, area in per_class.items():
            row = class_table.add_row().cells
            row[0].text = cls_name
            row[1].text = str(area)
            if total_defect > 0:
                row[2].text = f"{area / total_defect * 100:.1f}%"
            else:
                row[2].text = "0%"
        if total_defect == 0:
            doc.add_paragraph("当前置信度阈值下未检测到缺陷像素。")
    else:
        doc.add_paragraph("（无逐类数据）")

    # ---- 缺陷量化统计 ----
    doc.add_heading("三、缺陷量化统计", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "指标"
    hdr[1].text = "数值"

    rows_data = [
        ("缺陷类别", str(result_dict.get("defect_categories", ""))),
        ("缺陷面积（像素）", str(result_dict.get("defect_area", ""))),
        (
            "缺陷面积占比",
            f"{float(result_dict.get('defect_area_ratio', 0)) * 100:.4f}%",
        ),
        ("缺陷连通区域数量", str(result_dict.get("connected_components", ""))),
        ("最大缺陷区域面积（像素）", str(result_dict.get("max_defect_area", ""))),
        ("平均置信度", f"{float(result_dict.get('confidence_score', 0)):.4f}"),
        ("严重程度", str(result_dict.get("severity_level", ""))),
    ]
    for k, v in rows_data:
        row = table.add_row().cells
        row[0].text = k
        row[1].text = v

    # ---- 置信度分析 ----
    doc.add_heading("四、模型置信度分析", level=1)
    mean_conf_per_class = result_dict.get("mean_confidence_per_class") or {}
    if mean_conf_per_class:
        conf_table = doc.add_table(rows=1, cols=2)
        conf_table.style = "Table Grid"
        chdr = conf_table.rows[0].cells
        chdr[0].text = "类别"
        chdr[1].text = "平均置信度"
        for cls_name, mc in mean_conf_per_class.items():
            row = conf_table.add_row().cells
            row[0].text = cls_name
            row[1].text = f"{mc:.4f}"
    else:
        doc.add_paragraph("（无逐类置信度数据）")

    threshold = result_dict.get("confidence_threshold", 0.90)
    doc.add_paragraph(
        f"本次检测使用置信度阈值 {threshold}。"
        f"仅当模型对某像素的最高类别概率 ≥ {threshold} 时才将其计入缺陷区域。"
        f"阈值越高，误检率越低，但可能漏掉置信度较低的微弱缺陷。"
    )

    # ---- 运维建议 ----
    doc.add_heading("五、运维建议", level=1)
    doc.add_paragraph(str(result_dict.get("suggestion", "")))

    # ---- 检测结论 ----
    doc.add_heading("六、检测结论", level=1)
    sev = result_dict.get("severity_level", "正常")
    ratio = float(result_dict.get("defect_area_ratio", 0))
    defect_categories = result_dict.get("defect_categories", "")
    conclusion = (
        f"综合 PV-S3 语义分割结果（置信度阈值 {threshold}），"
        f"检测到缺陷类别：{defect_categories}，"
        f"缺陷面积占比为 {ratio * 100:.2f}%，"
        f"严重程度判定为「{sev}」。请结合现场工况参考运维建议执行后续动作。"
    )
    doc.add_paragraph(conclusion)

    doc.save(str(out_path))
    return str(out_path.resolve())
