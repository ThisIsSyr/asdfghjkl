# -*- coding: utf-8 -*-
"""AI Agent 增强报告：在规则模板 Word 上追加智能分析章节。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

from src.utils.config import PROJECT_ROOT


def append_ai_sections_to_docx(
    docx_path: str | Path,
    sections: dict[str, str],
    section_title: str = "七、智能分析报告（AI Agent 生成）",
) -> str:
    """
    向已有 Word 报告追加 AI 分析章节。

    Parameters
    ----------
    docx_path : 报告绝对或相对路径
    sections : {小标题: 正文段落}
    """
    path = Path(docx_path)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    if not path.is_file():
        raise FileNotFoundError(f"报告文件不存在: {path}")

    doc = Document(str(path))
    doc.add_page_break()
    heading = doc.add_heading(section_title, level=1)
    if heading.runs:
        heading.runs[0].font.size = Pt(14)

    disclaimer = doc.add_paragraph(
        "【说明】以下内容由大语言模型 Agent 根据 PV-S3 结构化检测结果自动生成，"
        "核心检测数据来自语义分割模型与后处理统计规则，仅供运维参考，"
        "重要决策请结合现场人工复核。"
    )
    disclaimer.runs[0].italic = True

    for subtitle, body in sections.items():
        if not body or not str(body).strip():
            continue
        doc.add_heading(subtitle, level=2)
        for para in str(body).split("\n"):
            text = para.strip()
            if text:
                doc.add_paragraph(text)

    doc.save(str(path))
    return str(path.resolve())


def build_agent_result_payload(
    report_path: str,
    agent_output: str,
    ctx_notes: list[str] | None = None,
) -> dict[str, Any]:
    return {
        "report_path": report_path,
        "agent_output": agent_output,
        "agent_notes": ctx_notes or [],
        "mode": "ai_agent",
    }
